"""A document's bytes as text, per format (DR-0483, DR-0484): what each record will hold, and the people the file names.

| Format | Gives |
|---|---|
| docx | one text: the body's paragraphs and table rows, then headers and footers, footnotes, endnotes, comments |
| pptx | a part per slide: its shapes and table rows, then its speaker notes |
| xlsx | a part per sheet: a paragraph per row from 1, each starting with its own row number, so paragraph n is row n |
| pdf | a part per page |
| md, txt | one text, as written |

Paragraphs are apart by one blank line, and one over 500 words is cut at sentence ends; a sheet's row never is. A file
past a cap raises `NotReadError`, and is kept out whole. Everything here runs in the extraction child (DR-0485), but md
and txt, which are only decoded.
"""

from __future__ import annotations

import contextlib
import io
import logging
import sys
import warnings
from collections.abc import Callable, Iterable, Iterator
from dataclasses import replace
from datetime import date, datetime, time, timedelta
from functools import partial
from importlib.metadata import version
from xml.etree.ElementTree import Element

from defusedxml import DefusedXmlException
from defusedxml.ElementTree import fromstring as defused_fromstring
from defusedxml.ElementTree import iterparse as defused_iterparse
from openpyxl import load_workbook
from pypdf import PasswordType, PdfReader

from docket.documents import text as texts
from docket.documents.limits import LIMITS, Limits, NotReadError, mib, over
from docket.documents.ooxml import UNSAFE_XML, Package, damaged, local, namespace
from docket.documents.parts import Extracted, Format, Part

NO_TEXT = "No text on this page: a scan or an image, and Docket doesn't run OCR."
NOTES = "Speaker notes:"

# WordprocessingML, PresentationML and DrawingML, in their transitional and strict namespaces.
W = frozenset(
    {"http://schemas.openxmlformats.org/wordprocessingml/2006/main", "http://purl.oclc.org/ooxml/wordprocessingml/main"}
)
P = frozenset(
    {"http://schemas.openxmlformats.org/presentationml/2006/main", "http://purl.oclc.org/ooxml/presentationml/main"}
)
A = frozenset({"http://schemas.openxmlformats.org/drawingml/2006/main", "http://purl.oclc.org/ooxml/drawingml/main"})
# Markup compatibility: a fallback repeats its choice for older readers, so it is never read.
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
DOCX_ENGINE = "docket-docx 1"
PPTX_ENGINE = "docket-pptx 1"
# Every parse openpyxl makes, through defusedxml with no DTD: openpyxl calls it with its defaults, which take one.
SAFE_FROMSTRING = partial(defused_fromstring, forbid_dtd=True)
SAFE_ITERPARSE = partial(defused_iterparse, forbid_dtd=True)


def extract(data: bytes, fmt: Format, limits: Limits = LIMITS) -> Extracted:
    """A file's text. Raises `NotReadError` for a file past a cap, damaged or protected, any error a reader meets
    named as damage by its type; a memory error goes on up, so the child ends."""
    if len(data) > limits.file_bytes:
        raise NotReadError("size", f"{mib(len(data))}, over the cap of {mib(limits.file_bytes)}")
    readers: dict[Format, Callable[[bytes, Limits], Extracted]] = {
        "pdf": read_pdf,
        "docx": read_docx,
        "xlsx": read_xlsx,
        "pptx": read_pptx,
        "md": read_text,
        "txt": read_text,
    }
    try:
        extracted = replace(readers[fmt](data, limits), format=fmt)
    except (NotReadError, MemoryError):
        raise
    except Exception as e:
        # defusedxml's refusal, ours or openpyxl's, which wraps it in an error of its own.
        if _refused_xml(e):
            raise NotReadError("xml", UNSAFE_XML) from None
        raise damaged(e) from None
    _check_text(extracted, limits)
    return extracted


def _refused_xml(error: BaseException) -> bool:
    """Whether defusedxml's refusal is the error or what caused it, a few links down at most."""
    chain = [error]
    while len(chain) < 8 and (link := chain[-1].__cause__ or chain[-1].__context__) is not None:
        chain.append(link)
    return any(isinstance(e, DefusedXmlException) for e in chain)


def harden_openpyxl() -> None:
    """Point the parse functions every loaded openpyxl module bound at import at the safe ones, and those of
    `openpyxl.xml.functions`, so a module imported later binds them too. `read_xlsx` runs it before openpyxl reads a
    workbook, so every part openpyxl parses refuses a DTD, whatever its name or content type says."""
    for name, module in list(sys.modules.items()):
        if not (name == "openpyxl" or name.startswith("openpyxl.")):
            continue
        for attribute, safe in (("fromstring", SAFE_FROMSTRING), ("iterparse", SAFE_ITERPARSE)):
            if attribute in vars(module):
                setattr(module, attribute, safe)


def _size(text: str) -> int:
    return len(text.encode("utf-8"))


def _check_text(extracted: Extracted, limits: Limits) -> None:
    for part in extracted.parts:
        if _size(part.text) > limits.part_text:
            raise NotReadError(
                "text", f"{part.title} holds {mib(_size(part.text))} of text, over the cap of {mib(limits.part_text)}"
            )
    total = _size(extracted.text) + sum(_size(p.text) for p in extracted.parts)
    if total > limits.file_text:
        raise NotReadError("text", f"{mib(total)} of text, over the cap of {mib(limits.file_text)}")


# -- md and txt -----------------------------------------------------------------------------------------------------


def read_text(data: bytes, limits: Limits) -> Extracted:
    text, note = texts.plain(data)
    return Extracted("txt", texts.TEXT_ENGINE, text=text, note=note)


# -- docx -----------------------------------------------------------------------------------------------------------


def _w(element: Element, name: str) -> bool:
    return local(element.tag) == name and namespace(element.tag) in W


# A content control or custom XML can wrap a table's rows or a row's cells; nothing else stands between them.
_WRAPPERS = frozenset({"sdt", "sdtContent", "customXml"})


def _own(parent: Element, name: str) -> Iterator[Element]:
    """A table's rows or a row's cells: its children of that name, and those under its wrappers, never deeper."""
    for child in parent:
        if _w(child, name):
            yield child
        elif namespace(child.tag) in W and local(child.tag) in _WRAPPERS:
            yield from _own(child, name)


class _Docx:
    """A WordprocessingML part's text: paragraphs and table rows in order, text boxes after the paragraph that holds
    them, tracked insertions in and deletions out, a field's result and never its code."""

    def blocks(self, container: Element) -> list[str]:
        out: list[str] = []
        for child in container:
            if namespace(child.tag) not in W:
                continue
            name = local(child.tag)
            if name == "p":
                boxes: list[str] = []
                out.append(self.paragraph(child, boxes))
                out.extend(boxes)
            elif name == "tbl":
                out.extend(self.rows(child))
            elif name in ("sdt", "sdtContent", "customXml", "smartTag", "ins", "moveTo"):
                out.extend(self.blocks(child))
        return out

    def rows(self, table: Element) -> list[str]:
        """A table's own rows, one line each: a table in a cell is read once, as that cell's text, so nesting never
        repeats it."""
        out: list[str] = []
        for row in _own(table, "tr"):
            cells = [texts.one_line(" ".join(self.blocks(cell))) for cell in _own(row, "tc")]
            out.append(" | ".join(cells))
        return out

    def paragraph(self, p: Element, boxes: list[str]) -> str:
        out: list[str] = []
        self._walk(p, out, boxes)
        return "".join(out)

    def _walk(self, element: Element, out: list[str], boxes: list[str]) -> None:
        for child in element:
            if child.tag == MC_FALLBACK:
                continue
            if namespace(child.tag) not in W:
                self._walk(child, out, boxes)
                continue
            name = local(child.tag)
            if name in ("del", "moveFrom", "delText", "instrText", "pPr", "rPr"):
                continue
            if name == "t":
                out.append(child.text or "")
            elif name == "tab":
                out.append("\t")
            elif name in ("br", "cr"):
                out.append("\n")
            elif name == "noBreakHyphen":
                out.append("-")
            elif name == "txbxContent":
                boxes.extend(self.blocks(child))
            elif name == "tbl":
                # A table inside a paragraph's run (rare, from some converters) reads as rows after it.
                boxes.extend(self.rows(child))
            else:
                self._walk(child, out, boxes)


def read_docx(data: bytes, limits: Limits) -> Extracted:
    package = Package(data, limits)
    main = package.main()
    reader = _Docx()
    root = package.xml(main)
    body = next((child for child in root if _w(child, "body")), None)
    blocks = reader.blocks(body) if body is not None else []
    extra: list[str] = []
    seen: set[str] = set()
    for kind in ("header", "footer"):
        for part in package.related(main, kind):
            for block in reader.blocks(package.xml(part)):
                if block.strip() and block not in seen:
                    seen.add(block)
                    extra.append(block)
    sections: list[str] = []
    if extra:
        sections += ["Headers and footers:", *extra]
    for kind, heading in (("footnotes", "Footnotes:"), ("endnotes", "Endnotes:")):
        notes = [
            block
            for part in package.related(main, kind)
            for note in package.xml(part)
            if note.get(_attr(note, "type")) not in ("separator", "continuationSeparator", "continuationNotice")
            for block in reader.blocks(note)
        ]
        if any(n.strip() for n in notes):
            sections += [heading, *notes]
    authors: list[str] = []
    comments: list[str] = []
    for part in package.related(main, "comments"):
        for comment in package.xml(part):
            author = comment.get(_attr(comment, "author"))
            if author and author.strip():
                name = " ".join(author.split())
                if name not in authors:
                    authors.append(name)
            comments.extend(reader.blocks(comment))
    if any(c.strip() for c in comments):
        sections += ["Comments:", *comments]
    creator, modified_by = package.core()
    paragraphs = texts.paragraphs([*blocks, *sections])
    return Extracted("docx", DOCX_ENGINE, text=texts.join(paragraphs), creator=creator, last_modified_by=modified_by,
                     authors=authors)  # fmt: skip


def _attr(element: Element, name: str) -> str:
    """An attribute in the element's own namespace, as WordprocessingML writes `w:type` and `w:author`."""
    ns = namespace(element.tag)
    return f"{{{ns}}}{name}" if ns else name


# -- pptx -----------------------------------------------------------------------------------------------------------


def _a(element: Element, name: str) -> bool:
    return local(element.tag) == name and namespace(element.tag) in A


def _p(element: Element, name: str) -> bool:
    return local(element.tag) == name and namespace(element.tag) in P


def _text_body(body: Element) -> str:
    """A text body's paragraphs, one line each, line breaks kept."""
    lines: list[str] = []
    for para in (e for e in body if _a(e, "p")):
        out: list[str] = []
        for run in para.iter():
            if _a(run, "t"):
                out.append(run.text or "")
            elif _a(run, "br"):
                out.append("\n")
        lines.append("".join(out))
    return "\n".join(lines)


def _shapes(tree: Element) -> Iterator[str]:
    """A shape tree's text in its order: a shape's text body as one paragraph, a table's rows one each, groups
    opened; pictures, charts and SmartArt give none."""
    for child in tree:
        if child.tag == MC_FALLBACK:
            continue
        if _p(child, "sp"):
            body = next((e for e in child if _p(e, "txBody")), None)
            if body is not None:
                yield _text_body(body)
        elif _p(child, "grpSp"):
            yield from _shapes(child)
        elif _p(child, "graphicFrame"):
            for table in (e for e in child.iter() if _a(e, "tbl")):
                for row in (r for r in table if _a(r, "tr")):
                    cells = [next((b for b in cell if _a(b, "txBody")), None) for cell in row if _a(cell, "tc")]
                    yield " | ".join(texts.one_line(_text_body(b)) if b is not None else "" for b in cells)
        else:
            # mc:AlternateContent and its choice hold shapes too.
            yield from _shapes(child)


def _notes(package: Package, slide: str) -> list[str]:
    """A slide's speaker notes: the body placeholder of its notes slide, never the slide image or its number."""
    out: list[str] = []
    for part in package.related(slide, "notesSlide"):
        tree = next((e for e in package.xml(part).iter() if _p(e, "spTree")), None)
        if tree is None:
            continue
        for shape in tree.iter():
            if not _p(shape, "sp"):
                continue
            placeholder = next((e for e in shape.iter() if _p(e, "ph")), None)
            body = next((e for e in shape if _p(e, "txBody")), None)
            if placeholder is not None and placeholder.get("type") == "body" and body is not None:
                out.append(_text_body(body))
    return out


def read_pptx(data: bytes, limits: Limits) -> Extracted:
    package = Package(data, limits)
    main = package.main()
    targets = {rel.id: rel.target for rel in package.rels(main)}
    ids = [e.get(R_ID, "") for e in package.xml(main).iter() if _p(e, "sldId")]
    order = [targets[i] for i in ids if targets.get(i) in package.names]
    if len(order) > limits.slides:
        raise over("slides", len(order), limits.slides)
    parts: list[Part] = []
    for n, slide in enumerate(order, start=1):
        root = package.xml(slide)
        hidden = root.get("show") == "0"
        tree = next((e for e in root.iter() if _p(e, "spTree")), None)
        blocks = list(_shapes(tree)) if tree is not None else []
        notes = _notes(package, slide)
        if any(note.strip() for note in notes):
            blocks += [NOTES, *notes]
        title = f"slide {n}" + (" (hidden)" if hidden else "")
        parts.append(Part("slide", n, title, texts.join(texts.paragraphs(blocks)), hidden))
    creator, modified_by = package.core()
    return Extracted("pptx", PPTX_ENGINE, parts=parts, creator=creator, last_modified_by=modified_by)


# -- xlsx -----------------------------------------------------------------------------------------------------------


def _cell(value: object) -> str:
    """A cell's saved value as text: dates as ISO dates, whole numbers without a decimal point."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return texts.one_line(str(value))


def read_xlsx(data: bytes, limits: Limits) -> Extracted:
    # Checked first: openpyxl opens the zip itself, and must meet an over-large or protected one only as refused.
    package = Package(data, limits)
    creator, modified_by = package.core()
    parts: list[Part] = []
    with warnings.catch_warnings():
        # openpyxl warns about features it drops (data validation, conditional formats): the text is what's kept.
        warnings.simplefilter("ignore")
        harden_openpyxl()
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            sheets = book.worksheets
            if len(sheets) > limits.sheets:
                raise over("sheets", len(sheets), limits.sheets)
            for n, sheet in enumerate(sheets, start=1):
                hidden = sheet.sheet_state != "visible"
                # The declared dimension is never trusted: a row is as wide as its own last cell.
                sheet.reset_dimensions()
                rows = _rows(sheet.iter_rows(values_only=True), n, limits)
                title = f"sheet {texts.one_line(sheet.title)}" + (" (hidden)" if hidden else "")
                parts.append(Part("sheet", n, title, texts.join(rows), hidden))
        finally:
            book.close()
    return Extracted("xlsx", f"openpyxl {version('openpyxl')}", parts=parts, creator=creator,
                     last_modified_by=modified_by)  # fmt: skip


def _rows(rows: Iterable[tuple[object, ...]], n: int, limits: Limits) -> list[str]:
    """A sheet's rows from 1 to the last with a value, row k as paragraph k. A cell past the column cap or a row past
    the row cap with a value keeps the file out."""
    out: list[str] = []
    pending: list[str] = []
    for number, row in enumerate(rows, start=1):
        if len(row) > limits.columns and any(v is not None for v in row[limits.columns :]):
            raise NotReadError("columns", f"sheet {n} has cells past column {limits.columns:,}")
        cells = [_cell(v) for v in row]
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            # Past the cap an empty row is never kept: a row with a value after it keeps the file out anyway.
            if number <= limits.rows:
                pending.append(f"{number}:")
            continue
        if number > limits.rows:
            raise NotReadError("rows", f"sheet {n} has rows past the cap of {limits.rows:,}")
        out.extend(pending)
        pending = []
        out.append(f"{number}: " + " | ".join(cells))
    return out


# -- pdf ------------------------------------------------------------------------------------------------------------


def read_pdf(data: bytes, limits: Limits) -> Extracted:
    # pypdf says what it repairs through logging and warnings: the text is what's kept.
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted and reader.decrypt("") == PasswordType.NOT_DECRYPTED:
            raise NotReadError("encrypted", "encrypted with a password")
        # The page tree as walked, never its own /Count: pypdf counts an encrypted file by /Count, which can lie.
        with contextlib.suppress(IndexError):
            reader.get_page(0)
        pages = list(reader.flattened_pages or [])
        if len(pages) > limits.pages:
            raise over("pages", len(pages), limits.pages)
        parts: list[Part] = []
        for n, page in enumerate(pages, start=1):
            paragraphs = texts.paragraphs(texts.blocks_of(page.extract_text() or ""))
            parts.append(Part("page", n, f"page {n}", texts.join(paragraphs) or NO_TEXT))
    return Extracted("pdf", f"pypdf {version('pypdf')}", parts=parts)
