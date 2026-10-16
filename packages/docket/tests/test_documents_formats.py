"""Each format's text as Docket keeps it (DR-0483): paragraphs apart by one blank line, a long one cut at sentence ends,
a sheet's row n as its paragraph n, and the people a file names. Every document is built here from invented words."""

from __future__ import annotations

import codecs
import datetime as dt

import pytest
from documents_kit import (
    Slide,
    docx,
    encrypted,
    pdf,
    pptx,
    sp,
    table,
    wp,
    wrow,
    xlsx,
)
from openpyxl import Workbook

from docket.documents import formats
from docket.documents import text as texts
from docket.documents.formats import NO_TEXT, NOTES, Part, extract, format_of

# -- text -------------------------------------------------------------------------------------------------------------


def test_text_is_nfc_with_lf_and_no_controls_but_the_tab() -> None:
    acute, line_separator, next_line = chr(0x301), chr(0x2028), chr(0x85)
    raw = f"Cafe{acute}\r\nline\rnext\x0bvert\x0cfeed{line_separator}sep\x00nul\x1fus\x7fdel{next_line}nel\tkept   \n"
    assert texts.clean(raw) == f"Caf{chr(0xE9)}\nline\nnext\nvert\nfeed\nsepnulusdel\nnel\tkept\n"
    assert texts.one_line("  two\nlines\t here ") == "two lines here"


def test_a_long_paragraph_is_cut_at_sentence_ends_and_a_long_sentence_at_a_space() -> None:
    sentence = " ".join(["word"] * 199) + " end."
    paragraph = " ".join([sentence] * 4)
    pieces = texts.split_long(paragraph)
    assert [len(p.split()) for p in pieces] == [400, 400]
    assert all(p.endswith("end.") for p in pieces)
    endless = " ".join(f"w{n}" for n in range(1, 1201))
    cut = texts.split_long(endless)
    assert [len(p.split()) for p in cut] == [500, 500, 200]
    assert cut[1].startswith("w501 ")
    assert texts.split_long("short one") == ["short one"]


def test_paragraphs_drop_blank_lines_inside_a_block_and_empty_blocks() -> None:
    long = " ".join(["word"] * 600)
    assert texts.paragraphs(["one\n\n  \ntwo", "", "  ", long]) == ["one\ntwo", *texts.split_long(long)]
    assert texts.paragraphs([long], split=False) == [long]
    assert texts.blocks_of("a\nb\n \n\t\n\nc") == ["a\nb", "c"]


@pytest.mark.parametrize(
    ("data", "text", "note"),
    [
        (codecs.BOM_UTF8 + "café".encode(), "café", None),
        ("café".encode("utf-16"), "café", None),
        ("café".encode(), "café", None),
        ("café".encode("cp1252"), "café", texts.WINDOWS_1252),
    ],
)
def test_md_and_txt_are_decoded_without_a_parser(data: bytes, text: str, note: str | None) -> None:
    assert texts.decode(data) == (text, note)


def test_a_txt_file_keeps_its_paragraphs_and_says_how_it_was_decoded() -> None:
    got = extract(f"Plan for Q3 caf{chr(0xE9)}\r\nsecond line\n\n\n\nNext part".encode("cp1252"), "txt")
    assert got.text == f"Plan for Q3 caf{chr(0xE9)}\nsecond line\n\nNext part"
    assert (got.format, got.note, got.parts) == ("txt", texts.WINDOWS_1252, [])
    assert extract(b"# Notes\n\nA line", "md").format == "md"


def test_format_by_extension() -> None:
    assert [format_of(n) for n in ("a.PDF", "b.docx", "c.Xlsx", "d.pptx", "e.md", "f.txt")] == [
        "pdf", "docx", "xlsx", "pptx", "md", "txt",
    ]  # fmt: skip
    assert format_of("g.doc") is None
    assert format_of("README") is None


# -- docx -------------------------------------------------------------------------------------------------------------

BODY = (
    wp("Northwind budget review") + "<w:p><w:r><w:t>Kept </w:t></w:r><w:ins><w:r><w:t>inserted </w:t></w:r></w:ins>"
    "<w:del><w:r><w:delText>deleted </w:delText></w:r></w:del>"
    "<w:moveFrom><w:r><w:t>moved away </w:t></w:r></w:moveFrom><w:r><w:t>end</w:t></w:r></w:p>"
    "<w:tbl><w:tblPr/>" + wrow("Item", "Qty") + wrow("Widget", "3") + "</w:tbl>"
    "<w:sdt><w:sdtPr/><w:sdtContent>" + wp("In a content control") + "</w:sdtContent></w:sdt>"
    "<w:p><w:r><mc:AlternateContent><mc:Choice><w:drawing><w:txbxContent>" + wp("Box text") + "</w:txbxContent>"
    "</w:drawing></mc:Choice><mc:Fallback><w:pict><w:txbxContent>" + wp("Box text") + "</w:txbxContent></w:pict>"
    "</mc:Fallback></mc:AlternateContent></w:r><w:r><w:t>Beside a box</w:t></w:r></w:p>"
    '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>PAGE</w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>7</w:t></w:r></w:p>'
    "<w:p><w:r><w:t>a</w:t><w:tab/><w:t>b</w:t><w:br/><w:t>c</w:t><w:noBreakHyphen/><w:t>d</w:t></w:r></w:p>"
    "<w:p><w:r><w:tbl>" + wrow("inline", "table") + "</w:tbl></w:r></w:p>"
    "<w:p/><w:sectPr/><other:x xmlns:other='urn:x'/>"
)


def test_a_docx_reads_as_it_would_print_with_its_changes_accepted() -> None:
    got = extract(
        docx(
            BODY,
            headers=[wp("Header line"), wp("Header line")],
            footnotes='<w:footnote w:type="separator" w:id="-1">' + wp("SEPARATOR") + "</w:footnote>"
            '<w:footnote w:id="1">' + wp("A footnote") + "</w:footnote>",
            endnotes='<w:endnote w:id="1">' + wp("An endnote") + "</w:endnote>",
            comments='<w:comment w:id="0" w:author="Ben  Sample">' + wp("Check the qty") + "</w:comment>"
            '<w:comment w:id="1" w:author="Ben Sample">' + wp("Again") + "</w:comment>"
            '<w:comment w:id="2" w:author=" ">' + wp("Nobody") + "</w:comment>",
            creator="Cleo Test",
            modified_by="Ben Sample",
        ),
        "docx",
    )
    assert got.text.split("\n\n") == [
        "Northwind budget review",
        "Kept inserted end",
        "Item | Qty",
        "Widget | 3",
        "In a content control",
        "Beside a box",
        "Box text",
        "7",
        "a\tb\nc-d",
        "inline | table",
        "Headers and footers:",
        "Header line",
        "Footnotes:",
        "A footnote",
        "Endnotes:",
        "An endnote",
        "Comments:",
        "Check the qty",
        "Again",
        "Nobody",
    ]
    assert (got.creator, got.last_modified_by, got.authors) == ("Cleo Test", "Ben Sample", ["Ben Sample"])
    assert (got.format, got.engine, got.parts) == ("docx", formats.DOCX_ENGINE, [])


def test_a_docx_without_a_body_notes_or_people_is_empty() -> None:
    got = extract(
        docx("", footnotes='<w:footnote w:type="separator">' + wp("x") + "</w:footnote>", comments=""), "docx"
    )
    assert (got.text, got.creator, got.authors) == ("", None, [])


def test_a_long_docx_paragraph_is_cut_at_sentence_ends() -> None:
    sentence = " ".join(["word"] * 299) + " stop."
    got = extract(docx(wp(f"{sentence} {sentence}")), "docx")
    assert got.text == f"{sentence}\n\n{sentence}"


def test_a_table_in_a_table_s_cell_is_read_once_as_that_cell_s_text() -> None:
    inner = "<w:tbl>" + wrow("Q1", "1,200") + wrow("Q2", "1,450") + "</w:tbl>"
    outer = "<w:tbl><w:tr><w:tc>" + wp("Budget") + "</w:tc><w:tc>" + inner + wp("") + "</w:tc></w:tr></w:tbl>"
    assert extract(docx(wp("Intro") + outer), "docx").text == "Intro\n\nBudget | Q1 | 1,200 Q2 | 1,450"


def test_rows_and_cells_inside_content_controls_are_the_table_s_own() -> None:
    cell = "<w:sdt><w:sdtContent><w:tc>" + wp("Harbour") + "</w:tc></w:sdtContent></w:sdt>"
    row = "<w:sdt><w:sdtContent><w:tr><w:tc>" + wp("Site") + "</w:tc>" + cell + "</w:tr></w:sdtContent></w:sdt>"
    wrapped = "<w:tbl><w:customXml>" + wrow("Date", "1 Oct") + "</w:customXml>" + row + "</w:tbl>"
    assert extract(docx(wrapped), "docx").text == "Date | 1 Oct\n\nSite | Harbour"


def test_deeply_nested_tables_stay_linear() -> None:
    body = wp("x")
    for _ in range(40):
        body = "<w:tbl><w:tr><w:tc>" + body + "</w:tc></w:tr></w:tbl>" + wp("")
    # Read whole at every level, 40 levels would be 2**40 copies of x.
    assert extract(docx(body), "docx").text == "x"


# -- pptx -------------------------------------------------------------------------------------------------------------


def test_a_deck_is_a_part_per_slide_in_the_deck_s_order_with_its_notes() -> None:
    group = "<p:grpSp>" + sp("In a group") + "</p:grpSp>"
    alternate = (
        "<mc:AlternateContent><mc:Choice>" + sp("Chosen") + "</mc:Choice><mc:Fallback>" + sp("Chosen")
        + "</mc:Fallback></mc:AlternateContent>"
    )  # fmt: skip
    picture = '<p:pic/><p:sp><p:nvSpPr><p:cNvPr id="2" name="no text"/></p:nvSpPr></p:sp>'
    breaks = "<p:sp><p:txBody><a:p><a:r><a:t>top</a:t></a:r><a:br/><a:r><a:t>bottom</a:t></a:r></a:p></p:txBody></p:sp>"
    empty_cell = "<p:graphicFrame><a:graphic><a:graphicData><a:tbl><a:tr><a:tc/><a:tc><a:txBody><a:p><a:r><a:t>x</a:t>"
    empty_cell += "</a:r></a:p></a:txBody></a:tc></a:tr></a:tbl></a:graphicData></a:graphic></p:graphicFrame>"
    deck = pptx(
        [
            Slide(sp("Second in the deck")),
            Slide(
                sp("Q3 plan", "for Northwind") + table(("Site", "Date"), ("Harbour", "1 Oct")) + group, notes="Say hi"
            ),
            Slide(sp("Hidden one") + alternate + picture + breaks + empty_cell, hidden=True, notes=""),
        ],
        creator="Cleo Test",
        order=[2, 1, 3],
    )
    got = extract(deck, "pptx")
    assert got.parts == [
        Part(
            "slide",
            1,
            "slide 1",
            f"Q3 plan\nfor Northwind\n\nSite | Date\n\nHarbour | 1 Oct\n\nIn a group\n\n{NOTES}\n\nSay hi",
        ),
        Part("slide", 2, "slide 2", "Second in the deck"),
        Part("slide", 3, "slide 3 (hidden)", "Hidden one\n\nChosen\n\ntop\nbottom\n\n | x", hidden=True),
    ]
    assert (got.creator, got.text, got.engine) == ("Cleo Test", "", formats.PPTX_ENGINE)


def test_a_slide_without_a_shape_tree_and_notes_without_a_body_give_no_text() -> None:
    from documents_kit import package, proot, rels

    deck = package(
        {
            "_rels/.rels": rels(("rId1", "officeDocument", "ppt/presentation.xml")),
            "ppt/presentation.xml": proot(
                "presentation",
                '<p:sldIdLst><p:sldId id="256" r:id="rId9"/><p:sldId id="257" r:id="rIdGone"/></p:sldIdLst>',
            ),
            "ppt/_rels/presentation.xml.rels": rels(("rId9", "slide", "slides/slide9.xml")),
            "ppt/slides/slide9.xml": proot("sld", "<p:cSld/>"),
            "ppt/slides/_rels/slide9.xml.rels": rels(("rIdN", "notesSlide", "../notesSlides/n.xml")),
            "ppt/notesSlides/n.xml": proot("notes", "<p:cSld/>"),
        }
    )
    assert extract(deck, "pptx").parts == [Part("slide", 1, "slide 1", "")]


# -- xlsx -------------------------------------------------------------------------------------------------------------


def _costs(book: Workbook) -> None:
    sheet = book.active
    assert sheet is not None
    sheet.title = "Costs"
    sheet["A1"], sheet["B1"] = "Item", "Qty"
    sheet["A3"], sheet["B3"], sheet["C3"] = "Widget", 3.0, 4.5
    sheet["A4"], sheet["B4"] = dt.datetime(2026, 10, 1), dt.datetime(2026, 10, 1, 9, 30)
    sheet["A5"], sheet["B5"], sheet["C5"] = dt.date(2026, 10, 2), dt.time(14, 5), dt.timedelta(hours=26)
    sheet["A6"], sheet["B6"], sheet["C6"] = "line one\nline two", True, False
    # A formula no app has calculated: openpyxl writes no saved value for it.
    sheet["A7"], sheet["B7"] = 12, "=B3*2"
    sheet["A9"] = 1e20
    old = book.create_sheet("Old  rates")
    old.sheet_state = "hidden"
    old["B2"] = "kept"


def test_a_workbook_is_a_part_per_sheet_with_row_n_as_paragraph_n() -> None:
    got = extract(xlsx(_costs, creator="Cleo Test"), "xlsx")
    assert got.parts == [
        Part(
            "sheet",
            1,
            "sheet Costs",
            "1: Item | Qty\n\n2:\n\n3: Widget | 3 | 4.5\n\n4: 2026-10-01 | 2026-10-01T09:30:00\n\n"
            "5: 2026-10-02 | 14:05:00 | 1 day, 2:00:00\n\n6: line one line two | TRUE | FALSE\n\n7: 12\n\n8:\n\n"
            "9: 1e+20",
        ),
        Part("sheet", 2, "sheet Old rates (hidden)", "1:\n\n2:  | kept", hidden=True),
    ]
    assert got.creator == "Cleo Test"
    assert got.engine.startswith("openpyxl ")
    paragraphs = got.parts[0].text.split("\n\n")
    assert all(p.split(":")[0] == str(n) for n, p in enumerate(paragraphs, start=1))


def test_a_saved_formula_value_is_read_and_a_row_is_never_cut() -> None:
    from documents_kit import replaced, sheet_xml

    def build(book: Workbook) -> None:
        sheet = book.active
        assert sheet is not None
        sheet["A1"] = "x"

    long = " ".join(["word"] * 700)
    rows = (
        '<row r="1"><c r="A1"><f>2*2</f><v>4.0</v></c></row>'
        f'<row r="3"><c r="B3" t="inlineStr"><is><t>{long}</t></is></c></row>'
    )
    data = replaced(xlsx(build), "xl/worksheets/sheet1.xml", sheet_xml(rows, dimension="A1:XFD1048576"))
    got = extract(data, "xlsx").parts[0].text
    assert got == f"1: 4\n\n2:\n\n3:  | {long}"


# -- pdf --------------------------------------------------------------------------------------------------------------


def test_a_pdf_is_a_part_per_page_and_a_page_without_text_says_so() -> None:
    got = extract(pdf([["Northwind invoice", "total (net) 10 000"], [], ["Page three"]]), "pdf")
    assert got.parts == [
        Part("page", 1, "page 1", "Northwind invoice\ntotal (net) 10 000"),
        Part("page", 2, "page 2", NO_TEXT),
        Part("page", 3, "page 3", "Page three"),
    ]
    assert got.engine.startswith("pypdf ")


def test_a_pdf_locked_by_an_owner_password_only_is_read() -> None:
    assert extract(encrypted(pdf([["Open to read"]]), ""), "pdf").parts[0].text == "Open to read"


@pytest.mark.parametrize("count", [b"1", b"5"])
def test_an_encrypted_pdf_s_pages_are_counted_by_its_page_tree_not_its_count(count: bytes) -> None:
    # pypdf counts an encrypted file by its /Count, which says what it likes: 1 hid two pages, 5 broke the read.
    data = encrypted(pdf([["page one"], ["page two"], ["page three"]]), "")
    assert data.count(b"/Count 3") == 1
    got = extract(data.replace(b"/Count 3", b"/Count " + count), "pdf")
    assert [p.text for p in got.parts] == ["page one", "page two", "page three"]


def test_a_pdf_without_pages_has_no_parts() -> None:
    assert extract(pdf([]), "pdf").parts == []
