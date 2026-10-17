"""`docket scan` on a temp folder (DR-0480, DR-0486, DR-0487): read-only, cheap on a second scan, a version only for
what changed, gone versions once, returns, moves, the guard, the gate's lists over the whole file, and nothing derived
from text before redaction in raw."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import stat
import sys
import threading
from collections import Counter
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
from docket_testkit import commit_config, git, host_repo, make_config
from documents_kit import docx, pdf, wp, xlsx
from documents_scan_kit import (
    MTIME,
    SOURCE,
    can_deny,
    cycle,
    folder_vault,
    junction,
    name,
    promote,
    put,
    raw,
    rid,
    run,
    staged,
    unlistable,
    versions,
)
from filelock import FileLock
from openpyxl import Workbook

from docket import folders, userconfig
from docket.documents import records
from docket.documents import scan as scanning
from docket.documents.extract import Done, Extractor, Failed, NotRead, Outcome
from docket.documents.limits import Limits
from docket.documents.parts import Extracted, Part
from docket.errors import DocketError
from docket.imports import notes as importing
from docket.vault.config import NeverStore, RedactionPattern, SourceConfig, VaultConfig
from docket.vault.gate import Gate
from docket.vault.records import RawRecord
from docket.vault.vault import Vault

NOW = datetime(2026, 10, 16, 8, 30, tzinfo=UTC)
LATER = datetime(2026, 10, 18, 9, 0, tzinfo=UTC)
# A value the password rule redacts: invented, and named so no code scanner takes the tests for storing a real one.
HIDDEN = "hunter2-" + "Example"


@pytest.fixture
def work(tmp_path: Path) -> Path:
    return tmp_path / "work"


@pytest.fixture
def vault(tmp_path: Path, config: VaultConfig, work: Path) -> Vault:
    return folder_vault(tmp_path, config, work)


Answers = dict[str, Callable[[Path], Outcome]]


class Asked:
    """What a scan asked its extractor for, by file name; a script answers some names instead of the real child."""

    def __init__(self, answers: Answers | None = None) -> None:
        self.answers = answers or {}
        self.asked: list[str] = []


class Counting(Extractor):
    """The real extractor, made only when the scan first needs it, as the scan makes its own."""

    def __init__(self, calls: Asked) -> None:
        super().__init__()
        self.calls = calls

    def extract(self, path: Path) -> Outcome:
        self.calls.asked.append(path.name)
        answer = self.calls.answers.get(path.name)
        return answer(path) if answer is not None else super().extract(path)


def counted(answers: Answers | None = None) -> tuple[Asked, Callable[[], Extractor]]:
    calls = Asked(answers)
    return calls, lambda: Counting(calls)


def really(path: Path) -> Outcome:
    with Extractor() as reader:
        return reader.extract(path)


@pytest.fixture
def reads(monkeypatch: pytest.MonkeyPatch) -> list[str]:
    """The names of the files the scan reads in its own process, to hash or to decode."""
    seen: list[str] = []
    real = scanning._read  # pyright: ignore[reportPrivateUsage]

    def reading(path: Path, *, keep: bool, cap: int) -> tuple[str, bytes, bytes | None]:
        seen.append(path.name)
        return real(path, keep=keep, cap=cap)

    monkeypatch.setattr(scanning, "_read", reading)
    return seen


def three_pages(second: str = "the second page") -> bytes:
    return pdf([["the first page"], [second], ["the third page"]])


def texts(vault: Vault, rid_: str) -> list[str]:
    return [r.text for r in versions(vault, rid_)]


def rewrite(vault: Vault, change: Callable[[VaultConfig], object]) -> Vault:
    config = VaultConfig.load(vault.root / "vault.yaml")
    change(config)
    (vault.root / "vault.yaml").write_text(config.dump(), encoding="utf-8", newline="\n")
    commit_config(vault.root)
    return Vault.open(vault.root)


# -- read-only, and cheap the second time ---------------------------------------------------------------------------


def test_a_scan_writes_nothing_in_a_folder_in_a_git_repo_and_keeps_its_ignore_rules(tmp_path: Path) -> None:
    top = host_repo(tmp_path / "host")
    work = top / "docs"
    put(work, "plan.md", "# Plan\n")
    put(work, "secret/kept-out.md", "# never read\n")
    (top / ".gitignore").write_text("docs/secret/\n", encoding="utf-8")
    put(work, "tracked.pdf", pdf([["a page"]]))
    git(top, "add", "docs/plan.md", ".gitignore")
    git(top, "commit", "-q", "-m", "docs")
    config = make_config("folder")
    config.sources[SOURCE] = SourceConfig(folder="docs")
    vault = Vault.init(top / "vault", config, "test-box")
    before = {p: (p.read_bytes(), p.stat().st_mtime_ns) for p in sorted(work.rglob("*")) if p.is_file()}
    index = (top / ".git" / "index").read_bytes()
    report = run(vault)
    assert report.outcomes == {"new": 2}
    assert {r.title for r in staged(vault) if r.kind == "document"} == {"plan.md", "tracked.pdf"}
    after = {p: (p.read_bytes(), p.stat().st_mtime_ns) for p in sorted(work.rglob("*")) if p.is_file()}
    assert after == before
    assert (top / ".git" / "index").read_bytes() == index


def test_a_second_scan_hashes_and_extracts_nothing_and_a_touched_file_makes_no_version(
    vault: Vault, work: Path, reads: list[str]
) -> None:
    put(work, "plan.md", "# Plan\n")
    put(work, "deck.pdf", three_pages())
    cycle(vault)
    reads.clear()
    made, factory = counted()
    report = cycle(vault, extractor=factory)
    assert (report.outcomes, reads, made.asked) == ({"same": 2}, [], [])
    put(work, "plan.md", "# Plan\n", mtime=MTIME + 60)
    put(work, "deck.pdf", three_pages(), mtime=MTIME + 60)
    made, factory = counted()
    report = run(vault, extractor=factory)
    assert report.outcomes == {"touched": 2}
    assert (sorted(reads), made.asked, staged(vault)) == (["deck.pdf", "plan.md"], [], [])
    reads.clear()
    assert run(vault).outcomes == {"same": 2}
    assert reads == []


def test_an_edit_to_one_page_versions_that_page_and_the_file_only(vault: Vault, work: Path) -> None:
    put(work, "deck.pdf", three_pages())
    cycle(vault)
    put(work, "deck.pdf", three_pages("the second page, edited"), mtime=MTIME + 60)
    report = run(vault)
    assert report.outcomes == {"changed": 1}
    assert [r.rid for r in staged(vault)] == [rid("deck.pdf"), f"{rid('deck.pdf')}/page2"]


def sheets(first_cell: str, other: str) -> Callable[[Workbook], None]:
    def build(book: Workbook) -> None:
        first = book.active
        assert first is not None
        first.append([first_cell])
        book.create_sheet("Other").append([other])

    return build


def test_a_credential_redacted_in_a_sheet_makes_no_version_when_another_sheet_changes(vault: Vault, work: Path) -> None:
    put(work, "access.xlsx", xlsx(sheets(f"password: {HIDDEN}", "first")))
    cycle(vault)
    sheet = f"{rid('access.xlsx')}/sheet1"
    assert texts(vault, sheet) == ["1: password: [redacted:password]"]
    put(work, "access.xlsx", xlsx(sheets(f"password: {HIDDEN}", "second")), mtime=MTIME + 60)
    run(vault)
    assert [r.rid for r in staged(vault)] == [rid("access.xlsx"), f"{rid('access.xlsx')}/sheet2"]


def test_a_new_redaction_rule_versions_the_parts_it_touches_once(vault: Vault, work: Path) -> None:
    put(work, "deck.pdf", pdf([["ticket PRJ-1234 is open"], ["nothing to hide"]]))
    cycle(vault)
    vault = rewrite(vault, lambda c: c.redaction.patterns.append(RedactionPattern(name="project", pattern=r"PRJ-\d+")))
    put(work, "deck.pdf", pdf([["ticket PRJ-1234 is open"], ["nothing to hide"], ["a new page"]]), mtime=MTIME + 60)
    cycle(vault)
    page1 = f"{rid('deck.pdf')}/page1"
    assert texts(vault, page1) == ["ticket PRJ-1234 is open", "ticket [redacted:project] is open"]
    assert len(versions(vault, f"{rid('deck.pdf')}/page2")) == 1
    put(work, "deck.pdf", pdf([["ticket PRJ-1234 is open"], ["nothing to hide"]]), mtime=MTIME + 120)
    run(vault)
    assert page1 not in {r.rid for r in staged(vault)}


def test_a_stat_that_matches_while_the_vault_s_newest_identity_differs_is_looked_at(vault: Vault, work: Path) -> None:
    put(work, "plan.md", "# Plan\n")
    cycle(vault)
    # Another machine's version of the file, as a pull would bring it.
    newest = versions(vault, rid("plan.md"))[-1]
    other = newest.model_copy(update={"text": "# Plan, theirs", "ts": LATER, "native": {
        **newest.native, records.TEXT_HASH: "0" * 64}})  # fmt: skip
    Gate(vault).stage([other])
    promote(vault)
    report = cycle(vault)
    assert report.outcomes == {"changed": 1}
    assert texts(vault, rid("plan.md"))[-1] == "# Plan"


def test_a_lost_cache_hashes_again_and_makes_no_version(vault: Vault, work: Path, reads: list[str]) -> None:
    put(work, "plan.md", "# Plan\n")
    put(work, "deck.pdf", three_pages())
    cycle(vault)
    (vault.home.private / "documents" / f"{SOURCE}.json").unlink()
    made, factory = counted()
    reads.clear()
    report = run(vault, extractor=factory)
    assert (report.outcomes, made.asked, staged(vault)) == ({"same": 2}, [], [])
    assert sorted(reads) == ["deck.pdf", "plan.md"]


@pytest.mark.parametrize(
    "text", ["not json", json.dumps({"schema": 0}), json.dumps({"schema": 1, "files": {"x": {"nope": 1}}})]
)
def test_a_cache_that_doesn_t_read_is_started_again(vault: Vault, work: Path, text: str) -> None:
    put(work, "plan.md", "# Plan\n")
    cycle(vault)
    (vault.home.private / "documents" / f"{SOURCE}.json").write_text(text, encoding="utf-8")
    assert run(vault).outcomes == {"same": 1}


def test_a_file_changed_while_it_was_read_drops_its_cache_entry(vault: Vault, work: Path) -> None:
    put(work, "deck.pdf", three_pages())

    def changing(path: Path) -> Outcome:
        outcome = really(path)
        put(work, "deck.pdf", three_pages("changed meanwhile"), mtime=MTIME + 60)
        return outcome

    _, factory = counted({"deck.pdf": changing})
    cycle(vault, extractor=factory)
    cache = json.loads((vault.home.private / "documents" / f"{SOURCE}.json").read_text(encoding="utf-8"))
    assert records.file_id("deck.pdf") not in cache["files"]
    report = run(vault)
    assert report.outcomes == {"changed": 1}


def test_a_file_deleted_while_it_was_read_is_looked_at_again(vault: Vault, work: Path) -> None:
    put(work, "deck.pdf", three_pages())

    def deleting(path: Path) -> Outcome:
        outcome = really(path)
        path.unlink()
        return outcome

    _, factory = counted({"deck.pdf": deleting})
    assert run(vault, extractor=factory).outcomes == {"new": 1}
    cache = json.loads((vault.home.private / "documents" / f"{SOURCE}.json").read_text(encoding="utf-8"))
    assert cache["files"] == {}


# -- gone, once; returns; moves -------------------------------------------------------------------------------------


def test_a_deleted_file_and_a_shrunk_pdf_get_gone_versions_once(vault: Vault, work: Path) -> None:
    put(work, "plan.md", "# Plan\n")
    put(work, "deck.pdf", three_pages())
    cycle(vault, now=NOW)
    (work / "plan.md").unlink()
    put(work, "deck.pdf", pdf([["the first page"]]), mtime=MTIME + 60)
    report = cycle(vault, now=NOW)
    assert (report.gone, report.outcomes) == (1, {"changed": 1})
    plan = versions(vault, rid("plan.md"))[-1]
    assert (plan.text, plan.ts, plan.title) == ("Not read from the folder since 2026-10-16.", NOW, "plan.md")
    assert plan.native == {"gone": "2026-10-16", "format": "md"}
    page3 = versions(vault, f"{rid('deck.pdf')}/page3")[-1]
    assert (page3.text, page3.native) == ("Not read from the file since 2026-10-16.", {"gone": "2026-10-16"})
    before = len(raw(vault))
    for later in (LATER, LATER.replace(day=19)):
        assert cycle(vault, now=later).gone == 0
    assert len(raw(vault)) == before


def test_a_file_restored_with_the_same_bytes_and_time_is_current_again(vault: Vault, work: Path) -> None:
    put(work, "deck.pdf", three_pages())
    put(work, "plan.md", "# Plan\n")
    cycle(vault, now=NOW)
    saved = (work / "deck.pdf").read_bytes()
    (work / "deck.pdf").unlink()
    (work / "plan.md").unlink()
    cycle(vault, now=NOW)
    put(work, "deck.pdf", saved)
    put(work, "plan.md", "# Plan\n")
    report = cycle(vault, now=LATER)
    assert report.outcomes == {"back": 2}
    for rid_ in (rid("deck.pdf"), f"{rid('deck.pdf')}/page2", rid("plan.md")):
        newest = versions(vault, rid_)[-1]
        assert not records.is_gone(newest)
        assert newest.ts == LATER
    assert texts(vault, rid("plan.md")) == ["# Plan", "Not read from the folder since 2026-10-16.", "# Plan"]


def test_a_lost_page_restored_is_current_again(vault: Vault, work: Path) -> None:
    put(work, "deck.pdf", three_pages())
    cycle(vault, now=NOW)
    put(work, "deck.pdf", pdf([["the first page"]]), mtime=MTIME + 60)
    cycle(vault, now=NOW)
    put(work, "deck.pdf", three_pages(), mtime=MTIME + 120)
    cycle(vault, now=LATER)
    page3 = versions(vault, f"{rid('deck.pdf')}/page3")
    assert [r.text for r in page3] == ["the third page", "Not read from the file since 2026-10-16.", "the third page"]
    assert page3[-1].ts == LATER
    assert len(versions(vault, f"{rid('deck.pdf')}/page1")) == 1


def test_a_move_names_where_it_went_and_where_it_came_from_and_a_move_back_is_current(vault: Vault, work: Path) -> None:
    put(work, "drafts/plan.md", "# Plan\n")
    put(work, "drafts/deck.pdf", three_pages())
    cycle(vault)
    for name_ in ("plan.md", "deck.pdf"):
        (work / "final").mkdir(exist_ok=True)
        (work / "drafts" / name_).rename(work / "final" / name_)
    report = cycle(vault, now=NOW)
    assert (report.outcomes, report.gone) == ({"moved": 2}, 2)
    for name_ in ("plan.md", "deck.pdf"):
        old, new = rid(f"drafts/{name_}"), rid(f"final/{name_}")
        assert versions(vault, old)[-1].native["moved_to"] == new
        assert versions(vault, new)[-1].native["moved_from"] == old
    (work / "final" / "plan.md").rename(work / "drafts" / "plan.md")
    report = cycle(vault, now=LATER)
    assert report.outcomes == {"back": 1, "same": 1}
    assert not records.is_gone(versions(vault, rid("drafts/plan.md"))[-1])


def test_an_excluded_file_is_gone_and_back_when_the_pattern_goes(vault: Vault, work: Path) -> None:
    put(work, "archive/old.md", "# Old\n")
    put(work, "Plan.Draft.md", "# Draft\n")
    cycle(vault)
    vault = rewrite(vault, lambda c: c.sources[SOURCE].exclude.extend(["archive/*", "*.draft.md"]))
    report = cycle(vault, now=NOW)
    assert (report.gone, report.skipped) == (2, {"excluded": 2})
    vault = rewrite(vault, lambda c: c.sources[SOURCE].exclude.clear())
    assert cycle(vault, now=LATER).outcomes == {"back": 2}


def test_the_guard_stages_nothing_when_most_files_went_until_asked(vault: Vault, work: Path) -> None:
    for n in range(24):
        put(work, f"note-{n:02}.md", f"# Note {n}\n")
    cycle(vault)
    for n in range(20):
        (work / f"note-{n:02}.md").unlink()
    put(work, "new.md", "# New\n")
    report = run(vault)
    assert (report.guarded, report.known, staged(vault)) == (20, 24, [])
    report = run(vault, accept_gone=True)
    assert (report.guarded, report.gone, report.outcomes) == (None, 20, {"same": 4, "new": 1})


def test_a_missing_folder_scans_nothing(vault: Vault, work: Path) -> None:
    put(work, "plan.md", "# Plan\n")
    cycle(vault)
    shutil.rmtree(work)
    report = run(vault)
    assert (report.missing, report.files, staged(vault)) == (True, 0, [])


# -- passing reasons ------------------------------------------------------------------------------------------------


def test_passing_reasons_never_make_a_file_gone(vault: Vault, work: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    for rel in ("locked.md", "cloud.md", "held.md", "slow.pdf", "deny.pdf"):
        put(work, rel, three_pages() if rel.endswith(".pdf") else f"# {rel}\n")
    cycle(vault)
    for rel in ("locked.md", "cloud.md", "held.md", "slow.pdf", "deny.pdf"):
        put(work, rel, three_pages("changed") if rel.endswith(".pdf") else f"# {rel}, changed\n", mtime=MTIME + 60)
    (vault.root / "state" / "holds").mkdir(parents=True, exist_ok=True)
    (vault.root / "state" / "holds" / f"{SOURCE}.txt").write_text(records.file_id("held.md") + "\n", encoding="utf-8")
    real_read = scanning._read  # pyright: ignore[reportPrivateUsage]
    read: list[str] = []

    def locked(path: Path, *, keep: bool, cap: int) -> tuple[str, bytes, bytes | None]:
        read.append(path.name)
        if path.name == "locked.md":
            raise PermissionError(13, "in use")
        return real_read(path, keep=keep, cap=cap)

    monkeypatch.setattr(scanning, "_read", locked)
    monkeypatch.setattr(folders, "placeholder", lambda info: info.st_size == (work / "cloud.md").stat().st_size)
    _, factory = counted({"slow.pdf": lambda _: Failed("stopped"), "deny.pdf": lambda _: Failed("unreadable")})
    report = cycle(vault, extractor=factory)
    assert report.outcomes == {"unreadable": 2, "placeholders": 1, "on hold": 1, "retry": 1}
    assert report.gone == 0
    # A held id and a placeholder aren't read at all.
    assert sorted(read) == ["deny.pdf", "locked.md", "slow.pdf"]
    for rel in ("locked.md", "cloud.md", "held.md", "slow.pdf", "deny.pdf"):
        assert not records.is_gone(versions(vault, rid(rel))[-1])
        assert len(versions(vault, rid(rel))) == 1


def test_a_read_that_stops_or_crashes_twice_on_the_same_bytes_is_a_record_saying_so(vault: Vault, work: Path) -> None:
    put(work, "slow.pdf", three_pages())
    put(work, "crash.pdf", three_pages())
    cycle(vault)
    put(work, "slow.pdf", three_pages("x"), mtime=MTIME + 60)
    put(work, "crash.pdf", three_pages("y"), mtime=MTIME + 60)
    answers: dict[str, Callable[[Path], Outcome]] = {
        "slow.pdf": lambda _: Failed("stopped"),
        "crash.pdf": lambda _: Failed("crashed"),
    }
    for expected in ({"retry": 2}, {"changed": 2}):
        _, factory = counted(answers)
        report = cycle(vault, extractor=factory, now=NOW)
        assert report.outcomes == expected
    assert texts(vault, rid("slow.pdf"))[-1] == "Not read: it took longer than the time limit twice."
    assert texts(vault, rid("crash.pdf"))[-1] == "Not read: it crashed the reader twice."
    assert versions(vault, f"{rid('crash.pdf')}/page2")[-1].native == {"gone": "2026-10-16"}
    assert sorted(report.not_read) == sorted([rid("slow.pdf"), rid("crash.pdf")])
    made, factory = counted(answers)
    assert run(vault, extractor=factory).outcomes == {"same": 2}
    assert made.asked == []


def test_a_pdf_past_the_page_cap_keeps_a_record_saying_so_and_its_pages_go(vault: Vault, work: Path) -> None:
    put(work, "deck.pdf", three_pages())
    cycle(vault, now=NOW)
    put(work, "deck.pdf", pdf([[]] * 2001), mtime=MTIME + 60)
    report = cycle(vault, now=NOW)
    assert report.not_read == [rid("deck.pdf")]
    file = versions(vault, rid("deck.pdf"))[-1]
    assert file.text == "Not read: 2,001 pages, over the cap of 2,000."
    assert (file.native["problem"], len(file.native[records.BYTES_HASH])) == ("pages", 64)
    assert all(records.is_gone(versions(vault, f"{rid('deck.pdf')}/page{n}")[-1]) for n in (1, 2, 3))


def test_md_and_txt_past_a_cap_hold_no_hash_and_are_done_while_their_stat_holds(
    vault: Vault, work: Path, reads: list[str]
) -> None:
    put(work, "big.md", "x" * 40)
    put(work, "wordy.txt", "word " * 5)
    limits = Limits(file_bytes=30, file_text=20)
    report = cycle(vault, limits=limits)
    assert sorted(report.not_read) == sorted([rid("big.md"), rid("wordy.txt")])
    assert texts(vault, rid("big.md")) == ["Not read: 1 MB, over the cap of 1 MB."]
    assert texts(vault, rid("wordy.txt")) == ["Not read: 1 MB of text, over the cap of 1 MB."]
    assert all(set(r.native) == {"format", "problem"} for r in raw(vault))
    reads.clear()
    assert run(vault, limits=limits).outcomes == {"same": 2}
    assert reads == []


def test_an_lfs_pointer_and_what_isn_t_a_document_are_left_out(
    vault: Vault, work: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    put(work, "big.pdf", "version https://git-lfs.github.com/spec/v1\noid sha256:" + "0" * 64 + "\nsize 9\n")
    put(work, "~$plan.docx", b"lock")
    put(work, "photo.PNG", b"png")
    put(work, "README", "no extension")
    put(work, "plan.md", "# Plan\n")
    put(work, "twin.md", "# Twin\n")
    real = folders.list_folder

    def doubled(*args: Any, **kwargs: Any) -> list[folders.Found]:
        found = real(*args, **kwargs)
        twin = next(f for f in found if f.rel == "twin.md")
        return [*found, twin]

    monkeypatch.setattr(folders, "list_folder", doubled)
    report = run(vault)
    assert report.outcomes == {"lfs pointers": 1, "new": 2}
    assert report.skipped == {"lock files": 1, "other .png": 1, "other (no extension)": 1, "same id": 1}


# -- the gate over the whole file ---------------------------------------------------------------------------------


def test_the_never_store_lists_keep_a_file_out_whole_and_take_back_what_the_vault_held(
    vault: Vault, work: Path, reads: list[str]
) -> None:
    put(work, "deck.pdf", pdf([["page"], ["page"], ["page"], ["page"], ["page zebrafinch"]]))
    put(work, "by-dora.docx", docx(wp("A brief."), creator="Dora Minta"))
    put(work, "commented.docx", docx(wp("Fine."), comments='<w:comment w:id="0" w:author="Dora Minta">' + wp("x")
                                     + "</w:comment>"))  # fmt: skip
    put(work, "later.md", "# Later\n")
    put(work, "zebrafinch notes.md", "# Never read\n")
    vault = rewrite(vault, lambda c: setattr(c, "never_store", NeverStore(people=["Dora Minta"],
                                                                         keywords=["zebrafinch"])))  # fmt: skip
    made, factory = counted()
    report = cycle(vault, extractor=factory)
    assert (report.outcomes, report.skipped) == ({"kept out": 3, "new": 1}, {"never-store": 1})
    assert "zebrafinch notes.md" not in made.asked + reads
    assert {r.title for r in raw(vault)} == {"later.md"}
    # A keyword added later takes back what the vault held, whole, and the file isn't read again while it holds.
    put(work, "later.md", "# Later, about zebrafinch\n", mtime=MTIME + 60)
    report = cycle(vault, now=NOW)
    assert (report.outcomes, report.gone) == ({"kept out": 4}, 1)
    assert texts(vault, rid("later.md"))[-1] == "Not read from the folder since 2026-10-16."
    reads.clear()
    assert run(vault).outcomes == {"kept out": 4}
    assert reads == []
    vault = rewrite(vault, lambda c: setattr(c, "never_store", NeverStore()))
    assert run(vault).outcomes == {"new": 4, "back": 1}


# -- where it runs, and how it is named ---------------------------------------------------------------------------


def test_another_machine_doesn_t_scan(vault: Vault, work: Path) -> None:
    vault = rewrite(vault, lambda c: setattr(c, "primary_machine", "other-box"))
    with pytest.raises(DocketError, match="scanned on the vault's primary machine, other-box, and this is test-box"):
        run(vault)


def test_what_isn_t_a_folder_source_isn_t_scanned(vault: Vault) -> None:
    with pytest.raises(DocketError, match=r"no source 'nope' in vault\.yaml; folder sources: documents"):
        scanning.scan(vault, "nope")
    with pytest.raises(DocketError, match="chat isn't a folder source"):
        scanning.scan(vault, "chat")
    vault = rewrite(vault, lambda c: setattr(c.sources[SOURCE], "enabled", False))
    with pytest.raises(DocketError, match="turned off"):
        run(vault)
    assert scanning.folder_sources(vault) == []
    with pytest.raises(DocketError, match=r"has no folder in vault\.yaml"):
        scanning.folder_of(vault, "chat", vault.config.sources["chat"])


def test_one_scan_at_a_time(vault: Vault, work: Path) -> None:
    put(work, "plan.md", "# Plan\n")
    folder = vault.home.private / "documents"
    folder.mkdir(parents=True, exist_ok=True)
    with FileLock(folder / scanning.LOCK_NAME), pytest.raises(DocketError, match="another docket scan"):
        run(vault)


def test_the_folder_is_found_under_home_in_the_repo_or_in_the_user_s_config(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    top = host_repo(tmp_path / "host")
    config = make_config("folder")
    config.sources[SOURCE] = SourceConfig(folder="docs/papers")
    vault = Vault.init(top / "vault", config, "test-box")
    put(top / "docs" / "papers", "plan.md", "# Plan\n")
    found = scanning.folder_of(vault, SOURCE, vault.config.sources[SOURCE])
    assert found == (top / "docs" / "papers").resolve()
    first = cycle(vault)
    # The folder moved inside its repo: the ids stay, since they come from the path inside it (DR-0482).
    (top / "docs" / "papers").rename(top / "docs" / "documents")
    vault = rewrite(vault, lambda c: setattr(c.sources[SOURCE], "folder", "docs/documents"))
    assert (first.outcomes, cycle(vault).outcomes) == ({"new": 1}, {"same": 1})
    home = tmp_path / "home"
    monkeypatch.setenv("HOME", str(home))
    monkeypatch.setenv("USERPROFILE", str(home))
    for value, under in (("~", home), ("~/Work", home / "Work")):
        entry = SourceConfig(folder=value)
        assert scanning.folder_of(vault, SOURCE, entry) == under.resolve()
    name(vault, tmp_path / "elsewhere")
    assert scanning.folder_of(vault, SOURCE, vault.config.sources[SOURCE]) == (tmp_path / "elsewhere").resolve()
    userconfig.set_section("folders", {str(vault.root): {SOURCE: str(top / "vault" / "raw")}})
    with pytest.raises(DocketError, match="inside the vault"):
        scanning.folder_of(vault, SOURCE, vault.config.sources[SOURCE])


def test_a_relative_folder_needs_the_vault_in_a_repo(vault: Vault) -> None:
    userconfig.set_section("folders", None)
    entry = SourceConfig(folder="docs")
    with pytest.raises(DocketError, match="relative to the repo that holds the vault"):
        scanning.folder_of(vault, SOURCE, entry)


def test_a_folder_that_is_a_link_or_a_junction_is_followed(tmp_path: Path, config: VaultConfig) -> None:
    real = tmp_path / "real"
    put(real, "plan.md", "# Plan\n")
    linked = tmp_path / "linked"
    if sys.platform == "win32":
        junction(real, linked)
    else:
        linked.symlink_to(real, target_is_directory=True)
    vault = folder_vault(tmp_path, config, linked)
    report = run(vault)
    assert (report.outcomes, report.folder) == ({"new": 1}, str(real.resolve()))
    assert [r.title for r in staged(vault)] == ["plan.md"]


# -- nothing derived from text before redaction in raw ----------------------------------------------------------


def test_raw_holds_no_hash_of_unredacted_text_and_no_md_or_txt_bytes_hash_or_size(vault: Vault, work: Path) -> None:
    said = f"password: {HIDDEN}"
    files = {
        "a.md": f"# A\n\n{said}\n".encode(),
        "b.txt": f"{said}\n".encode(),
        "c.docx": docx(wp(said)),
        "d.xlsx": xlsx(sheets(said, "other")),
    }
    for rel, data in files.items():
        put(work, rel, data)
    cycle(vault)
    stored = "".join(p.read_text(encoding="utf-8") for p in (vault.root / "raw").rglob("*.jsonl"))
    assert HIDDEN not in stored
    unredacted = [said, f"# A\n\n{said}", f"1: {said}"]
    hashes = [hashlib.sha256(t.encode()).hexdigest() for t in unredacted]
    hashes += [hashlib.sha256(files[rel]).hexdigest() for rel in ("a.md", "b.txt")]
    assert not any(h in stored for h in hashes)
    for rel in ("a.md", "b.txt"):
        native = versions(vault, rid(rel))[-1].native
        assert "size" not in native and records.BYTES_HASH not in native
    # The binary formats keep their bytes' hash, which their text can't rebuild.
    assert versions(vault, rid("c.docx"))[-1].native[records.BYTES_HASH] == hashlib.sha256(files["c.docx"]).hexdigest()


# -- the listing --------------------------------------------------------------------------------------------------


def test_placeholders_are_told_by_their_attributes_or_flags(monkeypatch: pytest.MonkeyPatch) -> None:
    class Info:
        def __init__(self, attributes: int = 0, flags: int = 0) -> None:
            self.st_file_attributes = attributes
            self.st_flags = flags

    for attributes in (0x400000, 0x40000, 0x1000):
        assert folders.placeholder(Info(attributes))  # type: ignore[arg-type]
    assert not folders.placeholder(Info(0x20))  # type: ignore[arg-type]
    monkeypatch.setattr(sys, "platform", "darwin")
    assert folders.placeholder(Info(flags=folders.SF_DATALESS))  # type: ignore[arg-type]
    monkeypatch.setattr(sys, "platform", "linux")
    assert not folders.placeholder(Info(flags=folders.SF_DATALESS))  # type: ignore[arg-type]


def test_what_isn_t_a_regular_file_or_is_gone_since_it_was_listed_is_left_out(tmp_path: Path) -> None:
    skipped: Counter[str] = Counter()
    fifo = os.stat_result((stat.S_IFIFO | 0o644, 0, 0, 1, 0, 0, 0, 0, 0, 0))
    assert folders._file(tmp_path / "pipe.md", "pipe.md", skipped, fifo) is None  # pyright: ignore[reportPrivateUsage]
    assert folders._file(tmp_path / "gone.md", "gone.md", skipped) is None  # pyright: ignore[reportPrivateUsage]
    assert skipped == {"not files": 1}


def test_a_staged_file_that_doesn_t_parse_is_passed_over(vault: Vault, work: Path) -> None:
    put(work, "plan.md", "# Plan\n")
    run(vault)
    (vault.root / "staging" / SOURCE / "broken.json").write_text("{", encoding="utf-8")
    assert run(vault).outcomes == {"same": 1}


def test_parts_already_staged_are_compared_with_what_staging_holds(vault: Vault, work: Path) -> None:
    put(work, "deck.pdf", three_pages())
    run(vault)
    put(work, "deck.pdf", three_pages("again"), mtime=MTIME + 60)
    run(vault)
    assert [r.rid for r in staged(vault)][-2:] == [rid("deck.pdf"), f"{rid('deck.pdf')}/page2"]


def test_a_part_extraction_names_its_hidden_parts(vault: Vault, work: Path) -> None:
    put(work, "deck.pdf", three_pages())
    answer = Done("0" * 64, 9, Extracted("pdf", "fake 1", parts=[Part("page", 1, "page 1", "x", hidden=True)]))
    _, factory = counted({"deck.pdf": lambda _: answer})
    run(vault, extractor=factory)
    assert staged(vault)[-1].native == {"hidden": True}


def test_an_md_or_txt_file_not_read_moves_by_its_path_alone(vault: Vault, work: Path) -> None:
    put(work, "big.md", "x" * 40)
    limits = Limits(file_bytes=30)
    cycle(vault, limits=limits)
    (work / "big.md").rename(work / "bigger.md")
    report = cycle(vault, limits=limits)
    assert (report.outcomes, report.gone) == ({"new": 1}, 1)
    assert "moved_to" not in versions(vault, rid("big.md"))[-1].native


def test_a_not_read_answer_from_the_reader_keeps_the_bytes_hash(vault: Vault, work: Path) -> None:
    put(work, "deck.pdf", three_pages())
    _, factory = counted({"deck.pdf": lambda _: NotRead("encrypted", "encrypted with a password", None, None)})
    run(vault, extractor=factory)
    record = staged(vault)[-1]
    assert record.text == "Not read: encrypted with a password."
    assert record.native[records.BYTES_HASH] == hashlib.sha256((work / "deck.pdf").read_bytes()).hexdigest()
    assert record.native["size"] == (work / "deck.pdf").stat().st_size


def test_records_go_to_the_gate_in_batches_and_an_interrupted_scan_goes_on_where_it_stopped(
    vault: Vault, work: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    for n in range(3):
        put(work, f"note-{n}.md", f"# Note {n}\n")
    monkeypatch.setattr(scanning, "BATCH_TEXT", 1)
    calls: list[int] = []
    real = Gate.stage

    def failing_second(self: Gate, batch: list[RawRecord]) -> Any:
        calls.append(len(batch))
        if len(calls) == 2:
            raise DocketError("stopped halfway")
        return real(self, batch)

    monkeypatch.setattr(Gate, "stage", failing_second)
    with pytest.raises(DocketError, match="stopped halfway"):
        run(vault)
    monkeypatch.setattr(Gate, "stage", real)
    report = run(vault)
    assert (calls, report.outcomes) == ([1, 1], {"same": 1, "new": 2})
    assert [r.title for r in staged(vault)] == ["note-0.md", "note-1.md", "note-2.md"]


def test_new_bytes_with_the_same_text_and_time_are_a_new_version_at_the_scan_s_time(vault: Vault, work: Path) -> None:
    put(work, "brief.docx", docx(wp("A brief."), creator="Ada Example"))
    cycle(vault, now=NOW)
    put(work, "brief.docx", docx(wp("A brief."), creator="Ben Sample"))
    assert cycle(vault, now=LATER).outcomes == {"changed": 1}
    newest = versions(vault, rid("brief.docx"))
    assert [(r.ts, r.native["creator"]) for r in newest] == [(newest[0].ts, "Ada Example"), (LATER, "Ben Sample")]
    assert cycle(vault).outcomes == {"same": 1}


# -- the review's findings ----------------------------------------------------------------------------------------

deny = pytest.mark.skipif(not can_deny(), reason="root lists every folder whatever its permissions")


@deny
def test_a_subfolder_that_can_t_be_listed_is_a_passing_reason_walked(vault: Vault, work: Path) -> None:
    put(work, "plan.md", "# Plan\n")
    for n in range(3):
        put(work, f"team/note-{n}.md", f"# Note {n}\n")
    cycle(vault)
    with unlistable(work / "team"):
        report = cycle(vault)
    assert (report.gone, report.unlisted) == (0, ["team"])
    assert report.outcomes == {"same": 1, "folders not listed": 1}
    assert not any(records.is_gone(r) for r in raw(vault))


@deny
def test_a_subfolder_git_can_t_open_is_a_passing_reason(tmp_path: Path) -> None:
    top = host_repo(tmp_path / "host")
    work = top / "docs"
    put(work, "plan.md", "# Plan\n")
    for n in range(3):
        put(work, f"team/note-{n}.md", f"# Note {n}\n")
    config = make_config("folder")
    config.sources[SOURCE] = SourceConfig(folder="docs")
    vault = Vault.init(top / "vault", config, "test-box")
    cycle(vault)
    with unlistable(work / "team"):
        report = cycle(vault)
    assert (report.gone, report.unlisted) == (0, ["team"])
    assert not any(records.is_gone(r) for r in raw(vault))


@deny
def test_a_bad_folder_can_be_excluded_and_a_never_store_one_isn_t_named(vault: Vault, work: Path) -> None:
    put(work, "plan.md", "# Plan\n")
    put(work, "team/note.md", "# Note\n")
    put(work, "zebrafinch/diary.md", "# Diary\n")
    cycle(vault)
    vault = rewrite(vault, lambda c: setattr(c, "never_store", NeverStore(keywords=["zebrafinch"])))
    with unlistable(work / "zebrafinch"):
        report = run(vault)
    assert (report.outcomes["folders not listed"], report.unlisted) == (1, [])
    vault = rewrite(vault, lambda c: c.sources[SOURCE].exclude.append("team/*"))
    with unlistable(work / "team"):
        report = cycle(vault)
    assert (report.gone, report.unlisted, report.skipped["excluded"]) == (1, [], 1)


@deny
def test_the_notes_import_skips_a_folder_it_can_t_list_and_says_so(tmp_path: Path) -> None:
    notes = tmp_path / "notes"
    put(notes, "a.md", "# a\n")
    put(notes, "locked/b.md", "# b\n")
    skipped: Counter[str] = Counter()
    unlisted: list[str] = []
    with unlistable(notes / "locked"):
        found = folders.list_folder(notes.resolve(), None, (tmp_path / "vault").resolve(), skipped, unlisted=unlisted)
        listed = importing.list_folder(notes.resolve(), "notes", None, (tmp_path / "vault").resolve(), skipped)
    assert ([f.rel for f in found], unlisted) == (["a.md"], ["locked"])
    assert ([item.old_path for item in listed], skipped) == (["notes/a.md"], {"folders not listed": 1})


def test_a_scan_reads_the_vault_under_its_lock_so_a_promotion_beside_it_never_doubles_a_version(
    vault: Vault, work: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    put(work, "deck.pdf", three_pages())
    put(work, "plan.md", "# Plan\n")
    cycle(vault)
    put(work, "deck.pdf", three_pages("changed"), mtime=MTIME + 60)
    put(work, "plan.md", "# Plan, changed\n", mtime=MTIME + 60)
    run(vault)
    real = scanning.Staging  # pyright: ignore[reportPrivateImportUsage]
    beside: list[threading.Thread] = []

    def promoting(*args: Any, **kwargs: Any) -> Any:
        # Another process promotes while the scan reads what the vault holds: it waits for the scan's lock.
        if not beside:
            other = Vault.open(vault.root)
            beside.append(threading.Thread(target=lambda: Gate(other).promote(SOURCE)))
            beside[0].start()
            beside[0].join(1.0)
        return real(*args, **kwargs)

    monkeypatch.setattr(scanning, "Staging", promoting)
    report = run(vault)
    beside[0].join(30)
    assert report.outcomes == {"same": 2}
    promote(vault)
    assert [r.text for r in versions(vault, rid("plan.md"))] == ["# Plan", "# Plan, changed"]


@pytest.mark.parametrize("limit", ["BATCH_FILES", "BATCH_SECONDS"])
def test_an_interrupted_first_scan_goes_on_where_it_stopped(
    vault: Vault, work: Path, monkeypatch: pytest.MonkeyPatch, limit: str
) -> None:
    for n in range(4):
        put(work, f"brief-{n}.docx", docx(wp(f"Brief {n}.")))
    monkeypatch.setattr(scanning, limit, 1 if limit == "BATCH_FILES" else 0.0)

    def interrupted(path: Path) -> Outcome:
        raise KeyboardInterrupt

    _, factory = counted({"brief-3.docx": interrupted})
    with pytest.raises(KeyboardInterrupt):
        run(vault, extractor=factory)
    made, factory = counted()
    report = run(vault, extractor=factory)
    assert (made.asked, report.outcomes) == (["brief-3.docx"], {"same": 3, "new": 1})


def test_a_folder_vault_yaml_names_never_leads_out_of_home_or_the_repo(vault: Vault) -> None:
    userconfig.set_section("folders", None)
    for value in ("~//Work", "~/../elsewhere"):
        entry = SourceConfig.model_construct(folder=value)
        with pytest.raises(DocketError, match="leads out of"):
            scanning.folder_of(vault, SOURCE, entry)


def test_a_file_gone_between_the_walk_s_listing_and_its_stat_is_left_out(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    class Vanished:
        name = "gone.md"
        path = str(tmp_path / "gone.md")

        def is_symlink(self) -> bool:
            return False

        def is_junction(self) -> bool:
            return False

        def is_dir(self, *, follow_symlinks: bool = True) -> bool:
            return False

        def stat(self, *, follow_symlinks: bool = True) -> os.stat_result:
            raise FileNotFoundError(2, "gone")

    monkeypatch.setattr(os, "scandir", lambda _: iter([Vanished()]))
    skipped: Counter[str] = Counter()
    assert folders.list_folder(tmp_path, None, tmp_path / "vault", skipped) == []
    assert skipped == {}


def test_a_folder_whose_git_can_t_be_looked_for_is_walked_into(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def denied(self: Path, *, follow_symlinks: bool = True) -> bool:
        raise PermissionError(13, "denied")

    monkeypatch.setattr(Path, "exists", denied)
    assert not folders._nested_repo(tmp_path)  # pyright: ignore[reportPrivateUsage]


@deny
def test_tracked_files_in_a_folder_that_can_t_be_read_never_turn_gone(tmp_path: Path) -> None:
    top = host_repo(tmp_path / "host")
    work = top / "docs"
    put(work, "plan.md", "# Plan\n")
    for n in range(3):
        put(work, f"team/note-{n}.md", f"# Note {n}\n")
    git(top, "add", "docs")
    git(top, "commit", "-q", "-m", "docs")
    config = make_config("folder")
    config.sources[SOURCE] = SourceConfig(folder="docs")
    vault = Vault.init(top / "vault", config, "test-box")
    cycle(vault)
    with unlistable(work / "team"):
        report = cycle(vault)
    assert report.gone == 0
    assert not any(records.is_gone(r) for r in raw(vault))


@pytest.mark.parametrize("in_git", [False, True], ids=["walked", "git"])
def test_a_file_whose_stat_is_refused_names_its_folder_once(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, in_git: bool
) -> None:
    top = host_repo(tmp_path / "host") if in_git else tmp_path / "plain"
    work = top / "docs"
    put(work, "plan.md", "# Plan\n")
    put(work, "team/a.md", "# A\n")
    put(work, "team/b.md", "# B\n")
    real = folders._file  # pyright: ignore[reportPrivateUsage]

    def refusing(path_: Path, rel: str, skipped: Counter[str], info: os.stat_result | None = None) -> Any:
        if rel.startswith("team/"):
            raise PermissionError(13, "denied")
        return real(path_, rel, skipped, info)

    monkeypatch.setattr(folders, "_file", refusing)
    unlisted: list[str] = []
    repo = top.resolve() if in_git else None
    found = folders.list_folder(work.resolve(), repo, (tmp_path / "vault").resolve(), Counter(), unlisted=unlisted)
    assert ([f.rel for f in found], unlisted) == (["plan.md"], ["team"])


def test_tracked_files_whose_lstat_is_refused_keep_their_last_version(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """What POSIX does to the files of a folder without permissions, played on every system: each lstat is refused,
    the link check's included."""
    top = host_repo(tmp_path / "host")
    work = top / "docs"
    put(work, "plan.md", "# Plan\n")
    for n in range(2):
        put(work, f"team/note-{n}.md", f"# Note {n}\n")
    git(top, "add", "docs")
    git(top, "commit", "-q", "-m", "docs")
    config = make_config("folder")
    config.sources[SOURCE] = SourceConfig(folder="docs")
    vault = Vault.init(top / "vault", config, "test-box")
    cycle(vault)
    real = Path.lstat
    team = (work / "team").resolve()

    def refusing(self: Path) -> os.stat_result:
        if team in self.resolve().parents:
            raise PermissionError(13, "denied")
        return real(self)

    monkeypatch.setattr(Path, "lstat", refusing)
    report = cycle(vault)
    assert (report.gone, report.unlisted, report.outcomes) == (0, ["team"], {"same": 1, "folders not listed": 1})
